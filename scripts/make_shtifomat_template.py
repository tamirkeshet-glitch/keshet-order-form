# בונה את order-template-shtifomat.xlsx - תבנית רכבי שטיפומט להורדה מהטופס.
# הכותרות והערכים חייבים להיות זהים ל-SHTIFO_EXCEL_COLUMNS ול-SHTIFO_VTYPES ב-order.js.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["מס׳ רכב", "סוג הרכב", "כמות שטיפות בחודש"]
VTYPES = ["פרטי", "מסחרי", "מסחרי גדול", "משאית", "רכב משא"]
QTY = ["1", "2", "3", "4", "5"]
ROWS = 300

wb = Workbook()
ws = wb.active
ws.title = "שטיפומט"
ws.sheet_view.rightToLeft = True
red = PatternFill("solid", fgColor="E30613")
for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h + " *")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = red
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[cell.column_letter].width = 22
ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"
# מספר רכב כטקסט - אחרת אקסל מוחק אפס מוביל
for r in range(2, ROWS + 2):
    ws[f"A{r}"].number_format = "@"

lists = wb.create_sheet("רשימות")
lists.sheet_state = "hidden"
for c, values in enumerate([VTYPES, QTY], 1):
    for r, v in enumerate(values, 1):
        lists.cell(row=r, column=c, value=v)

def dropdown(col, list_col, n):
    dv = DataValidation(type="list", formula1=f"='רשימות'!${list_col}$1:${list_col}${n}",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="ערך לא תקין", error="יש לבחור ערך מהרשימה")
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{ROWS + 1}")

dropdown("B", "A", len(VTYPES))
dropdown("C", "B", len(QTY))

help_ws = wb.create_sheet("הוראות", 0)
help_ws.sheet_view.rightToLeft = True
help_ws.column_dimensions["A"].width = 100
lines = [
    "תבנית רכבים לשטיפומט - קשת יזמות עסקית / סונול",
    "",
    "1. ממלאים שורה לכל רכב בגיליון «שטיפומט». כל העמודות חובה.",
    "2. סוג הרכב וכמות השטיפות בחודש (1 עד 5) נבחרים מהרשימה הנפתחת.",
    "3. אין לשנות את שורת הכותרות.",
    "4. שומרים את הקובץ ומעלים אותו בטופס, במסך «עדכון שטיפומט לרכב».",
    "",
    "כל הרכבים ייכנסו לטופס שטיפומט אחד לחתימה.",
    "מחיר שטיפה (לא כולל מע״מ): פרטי 39 ₪ · מסחרי 45 ₪ · מסחרי גדול 69 ₪ · משאית 83 ₪ · רכב משא 128 ₪.",
]
for r, t in enumerate(lines, 1):
    help_ws.cell(row=r, column=1, value=t).font = Font(bold=(r == 1), size=13 if r == 1 else 11)
wb.active = 1
wb.save("order-template-shtifomat.xlsx")
print("ok")
