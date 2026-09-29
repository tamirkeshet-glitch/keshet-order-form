# בונה את order-template.xlsx - תבנית ההזמנה להורדה מהטופס.
# הכותרות והערכים ברשימות חייבים להיות זהים ל-order.js (EXCEL_COLUMNS, AMTSEI, FUEL, VTYPES),
# כי הטופס קורא את הקובץ לפי שם הכותרת ומשווה ערכים לרשימות האלה.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["סוג אמצעי תדלוק", "מס׳ רכב", "מס׳ טלפון נהג", "סוג דלק", "סוג הרכב",
           "הגבלה בליטרים ליום", "הגבלה בליטרים לחודש", "דגם רכב", "שנת יצור",
           "שם נהג", "קוד / שם מחלקה", "שטיפומט"]
REQUIRED = {0, 1, 3, 4}
AMTSEI = ["כרטיס רכב", "דלקן א'", "דלקן ב (רושם ק\"מ)", "דלקן אוריאה", "כרטיס אוריאה"]
FUEL = ["בנזין 95", "בנזין 98", "גולדיזל (סולר)", "אוריאה", "גולדיזל + אוריאה"]
VTYPES = ["פרטי", "מסחרי", "משאית", "אוטובוס", "אופנוע", "אחר"]
SHTIFO = ["לא", "1", "2", "3", "4", "5"]
ROWS = 300

wb = Workbook()
ws = wb.active
ws.title = "הזמנה"
ws.sheet_view.rightToLeft = True
red = PatternFill("solid", fgColor="E30613")
dark = PatternFill("solid", fgColor="1A2530")
for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h + (" *" if c - 1 in REQUIRED else ""))
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = red if c - 1 in REQUIRED else dark
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[cell.column_letter].width = 18
ws.row_dimensions[1].height = 32
ws.freeze_panes = "A2"
# מספר רכב, טלפון ושנה כטקסט - אחרת אקסל מוחק אפס מוביל
for col in "BCI":
    for r in range(2, ROWS + 2):
        ws[f"{col}{r}"].number_format = "@"

lists = wb.create_sheet("רשימות")
lists.sheet_state = "hidden"
for c, values in enumerate([AMTSEI, FUEL, VTYPES, SHTIFO], 1):
    for r, v in enumerate(values, 1):
        lists.cell(row=r, column=c, value=v)

def dropdown(col, list_col, n):
    dv = DataValidation(type="list", formula1=f"='רשימות'!${list_col}$1:${list_col}${n}",
                        allow_blank=True, showErrorMessage=True,
                        errorTitle="ערך לא תקין", error="יש לבחור ערך מהרשימה")
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{ROWS + 1}")

dropdown("A", "A", len(AMTSEI))
dropdown("D", "B", len(FUEL))
dropdown("E", "C", len(VTYPES))
dropdown("L", "D", len(SHTIFO))

help_ws = wb.create_sheet("הוראות", 0)
help_ws.sheet_view.rightToLeft = True
help_ws.column_dimensions["A"].width = 100
lines = [
    "תבנית הזמנת דלקן / כרטיס רכב - קשת יזמות עסקית / סונול",
    "",
    "1. ממלאים שורה לכל רכב בגיליון «הזמנה». עמודות באדום (*) הן חובה.",
    "2. בעמודות עם רשימה נפתחת בוחרים ערך מהרשימה בלבד.",
    "3. שטיפומט: «לא», או מספר השטיפות בחודש (1 עד 5). ריק = «לא».",
    "4. אין לשנות את שורת הכותרות ואין להוסיף עמודות.",
    "5. שומרים את הקובץ ומעלים אותו בטופס, במסך «דלקן / כרטיס רכב».",
    "",
    "הזמנה עם יותר מ-10 רכבים, או עם כמה סוגי אמצעי תדלוק,",
    "תפוצל אוטומטית לכמה טפסים. כל טופס יישלח לחתימה בנפרד.",
]
for r, t in enumerate(lines, 1):
    help_ws.cell(row=r, column=1, value=t).font = Font(bold=(r == 1), size=13 if r == 1 else 11)
wb.active = 1
wb.save("order-template.xlsx")
print("ok")
